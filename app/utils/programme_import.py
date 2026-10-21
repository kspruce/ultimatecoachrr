"""Import a training programme from a spreadsheet, and export one back.

The format is documented in docs/PROGRAMME_IMPORT.md; the export writes the
same format, so an exported programme can be imported again. Two sheets
(or two CSV files):

- the programme: one row per exercise (or per session, or per rest week),
  with the track, week and session it belongs to;
- the tests (optional): one row per performance test, with the weeks of its
  baseline and retests.

parse() never writes anything: it returns what it would build, with every
problem tied to a row number so the coach can fix the sheet and try again.
build() then creates a new programme (a draft) from a clean parse.
"""
import csv
import hashlib
import io
import json

from app import db
from app.models.programme import (
    AGGREGATIONS, LOG_FIELDS, MEASUREMENTS, REQUIREMENTS, SESSION_TYPES, WEEK_KINDS,
    PerformanceTest, Programme, ProgrammeExercise, ProgrammeMilestone, ProgrammeSession, ProgrammeTrack,
    ProgrammeVersion, ProgrammeWeek,
)
from app.utils.media_embed import is_web_url

PROGRAMME_COLUMNS = [
    'track', 'track_first_week', 'track_start_offset', 'week', 'week_title', 'week_type', 'week_objective',
    'week_skills', 'session', 'session_type', 'requirement', 'alternative_group', 'session_minutes',
    'session_focus', 'warmup', 'cooldown', 'coaching_cues', 'session_video', 'log_fields',
    'exercise', 'block', 'sets', 'reps', 'distance_m', 'duration_s', 'rest', 'intensity', 'tempo',
    'instructions', 'exercise_video',
]
REQUIRED_COLUMNS = ['track', 'week', 'session']
TEST_COLUMNS = ['test', 'unit', 'measurement', 'attempts', 'aggregation', 'better', 'min', 'max',
                'description', 'instructions', 'milestone_track', 'baseline_week', 'retest_weeks']
TEST_REQUIRED = ['test', 'unit']

SESSION_TYPE_KEYS = {k for k, _ in SESSION_TYPES}
REQUIREMENT_KEYS = {k for k, _ in REQUIREMENTS}
WEEK_KIND_KEYS = {k for k, _ in WEEK_KINDS}
MEASUREMENT_KEYS = {k for k, _ in MEASUREMENTS}
AGGREGATION_KEYS = {k for k, _ in AGGREGATIONS}
LOG_FIELD_KEYS = {k for k, _ in LOG_FIELDS}
LOG_FIELD_BY_LABEL = {label.lower(): k for k, label in LOG_FIELDS}

MAX_ROWS = 5000


class Problems:
    def __init__(self):
        self.errors, self.warnings = [], []

    def error(self, row, msg):
        self.errors.append((row, msg))

    def warn(self, row, msg):
        self.warnings.append((row, msg))


# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------

def rows_from_csv(text):
    text = (text or '').lstrip('﻿')
    if not text.strip():
        return []
    return [row for row in csv.reader(io.StringIO(text))]


def rows_from_xlsx(data):
    """{sheet name (lower case): rows} from an .xlsx file's bytes."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(['' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                         for v in row])
        sheets[ws.title.strip().lower()] = rows
    return sheets


def rows_to_csv(rows):
    out = io.StringIO()
    writer = csv.writer(out, lineterminator='\n')
    for row in rows:
        writer.writerow(row)
    return out.getvalue()


def _records(rows, known, required, problems, what):
    """[(row number, {column: value})] with headers normalised; reports
    missing and unknown columns."""
    rows = [r for r in rows if any((c or '').strip() for c in r)]
    if not rows:
        return []
    header = [(h or '').strip().lower().replace(' ', '_') for h in rows[0]]
    missing = [c for c in required if c not in header]
    if missing:
        problems.error(1, f'The {what} sheet is missing the column(s): {", ".join(missing)}.')
        return []
    unknown = [h for h in header if h and h not in known]
    if unknown:
        problems.warn(1, f'Ignored unknown {what} column(s): {", ".join(unknown)}.')
    if len(rows) > MAX_ROWS:
        problems.error(1, f'The {what} sheet has more than {MAX_ROWS} rows.')
        return []
    out = []
    for i, row in enumerate(rows[1:], start=2):
        record = {}
        for col, value in zip(header, row):
            if col in known:
                record[col] = (value or '').strip()
        out.append((i, record))
    return out


def content_hash(programme_rows, test_rows):
    def norm(rows):
        return [[(c or '').strip() for c in r] for r in rows if any((c or '').strip() for c in r)]
    payload = json.dumps([norm(programme_rows), norm(test_rows or [])], separators=(',', ':'))
    return hashlib.sha256(payload.encode()).hexdigest()


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------

def _int(value, row, problems, label, minimum=None):
    if value in (None, ''):
        return None
    try:
        number = int(float(value))
        if float(value) != number:
            raise ValueError
    except ValueError:
        problems.error(row, f'{label} "{value}" must be a whole number.')
        return None
    if minimum is not None and number < minimum:
        problems.error(row, f'{label} can\'t be below {minimum}.')
        return None
    return number


def _float(value, row, problems, label, minimum=None):
    if value in (None, ''):
        return None
    try:
        number = float(value)
    except ValueError:
        problems.error(row, f'{label} "{value}" must be a number.')
        return None
    if number != number or (minimum is not None and number < minimum):
        problems.error(row, f'{label} can\'t be below {minimum}.')
        return None
    return number


def _set_once(target, key, value, row, problems, label):
    """Week and session details can be given on any of their rows; the first
    wins, and a different value later is reported."""
    if value in (None, ''):
        return
    if target.get(key) in (None, ''):
        target[key] = value
    elif target[key] != value:
        problems.warn(row, f'{label} differs from an earlier row ("{target[key]}"); kept the earlier one.')


def _log_fields(value, row, problems):
    if not value:
        return None
    keys = []
    for part in value.split(','):
        part = part.strip().lower()
        if not part:
            continue
        key = part if part in LOG_FIELD_KEYS else LOG_FIELD_BY_LABEL.get(part)
        if key is None:
            problems.error(row, f'Unknown log field "{part}" (use {", ".join(sorted(LOG_FIELD_KEYS))}).')
            continue
        keys.append(key)
    return keys


def parse(programme_rows, test_rows=None):
    """What an import would create, plus the problems found."""
    problems = Problems()
    records = _records(programme_rows, PROGRAMME_COLUMNS, REQUIRED_COLUMNS, problems, 'programme')
    tracks = {}
    order = []
    last = {'track': None, 'week': None, 'session': None}
    for row, rec in records:
        track_name = rec.get('track') or last['track']
        if not track_name:
            problems.error(row, 'No track: fill in the track column.')
            continue
        if rec.get('track') and rec['track'] != last['track']:
            last['week'] = last['session'] = None
        week_raw = rec.get('week') or last['week']
        week_no = _int(week_raw, row, problems, 'Week')
        if week_no is None:
            if not week_raw:
                problems.error(row, 'No week number.')
            continue
        if str(week_no) != str(last['week']):
            last['session'] = None
        last['track'], last['week'] = track_name, str(week_no)

        track = tracks.get(track_name.lower())
        if track is None:
            track = tracks[track_name.lower()] = {'name': track_name, 'first_week': None, 'offset': None,
                                                 'weeks': {}, 'row': row}
            order.append(track)
        first = _int(rec.get('track_first_week'), row, problems, 'track_first_week')
        offset = _int(rec.get('track_start_offset'), row, problems, 'track_start_offset', minimum=0)
        _set_once(track, 'first_week', first, row, problems, 'track_first_week')
        _set_once(track, 'offset', offset, row, problems, 'track_start_offset')

        week = track['weeks'].setdefault(week_no, {'number': week_no, 'sessions': [], 'row': row})
        kind = (rec.get('week_type') or '').lower()
        if kind and kind not in WEEK_KIND_KEYS:
            problems.error(row, f'week_type "{kind}" must be one of {", ".join(sorted(WEEK_KIND_KEYS))}.')
            kind = ''
        _set_once(week, 'kind', kind, row, problems, 'week_type')
        _set_once(week, 'title', rec.get('week_title'), row, problems, 'week_title')
        _set_once(week, 'objective', rec.get('week_objective'), row, problems, 'week_objective')
        _set_once(week, 'skills', rec.get('week_skills'), row, problems, 'week_skills')

        if not rec.get('exercise') and any(rec.get(c) for c in ('sets', 'reps', 'distance_m', 'duration_s', 'rest',
                                                                  'intensity', 'tempo', 'instructions')):
            problems.error(row, 'Sets, reps, distance, rest or instructions given without an exercise name.')
        session_title = rec.get('session') or (last['session'] if rec.get('exercise') else None)
        if not session_title:
            if rec.get('exercise'):
                problems.error(row, f'Exercise "{rec["exercise"]}" has no session.')
            continue
        last['session'] = session_title
        session = next((s for s in week['sessions'] if s['title'].lower() == session_title.lower()), None)
        if session is None:
            session = {'title': session_title, 'exercises': [], 'row': row}
            week['sessions'].append(session)
        stype = (rec.get('session_type') or '').strip()
        if stype:
            if stype.lower() in SESSION_TYPE_KEYS:
                _set_once(session, 'session_type', stype.lower(), row, problems, 'session_type')
            else:
                _set_once(session, 'session_type', 'custom', row, problems, 'session_type')
                _set_once(session, 'custom_type', stype, row, problems, 'session_type')
        req = (rec.get('requirement') or '').lower()
        if req and req not in REQUIREMENT_KEYS:
            problems.error(row, f'requirement "{req}" must be required, optional or alternative.')
            req = ''
        _set_once(session, 'requirement', req, row, problems, 'requirement')
        _set_once(session, 'alternative_group', rec.get('alternative_group'), row, problems, 'alternative_group')
        _set_once(session, 'duration_minutes', _int(rec.get('session_minutes'), row, problems, 'session_minutes', 0),
                  row, problems, 'session_minutes')
        for col, key in (('session_focus', 'focus'), ('warmup', 'warmup'), ('cooldown', 'cooldown'),
                         ('coaching_cues', 'coaching_cues')):
            _set_once(session, key, rec.get(col), row, problems, col)
        video = rec.get('session_video')
        if video and not is_web_url(video):
            problems.error(row, f'session_video "{video}" isn\'t a web link.')
            video = None
        _set_once(session, 'video_url', video, row, problems, 'session_video')
        fields = _log_fields(rec.get('log_fields'), row, problems)
        if fields is not None and 'log_fields' not in session:
            session['log_fields'] = fields

        if rec.get('exercise'):
            ex_video = rec.get('exercise_video')
            if ex_video and not is_web_url(ex_video):
                problems.error(row, f'exercise_video "{ex_video}" isn\'t a web link.')
                ex_video = None
            session['exercises'].append({
                'name': rec['exercise'], 'block': rec.get('block') or None, 'sets': rec.get('sets') or None,
                'reps': rec.get('reps') or None,
                'distance_m': _float(rec.get('distance_m'), row, problems, 'distance_m', 0),
                'duration_seconds': _int(rec.get('duration_s'), row, problems, 'duration_s', 0),
                'rest': rec.get('rest') or None, 'intensity': rec.get('intensity') or None,
                'tempo': rec.get('tempo') or None, 'instructions': rec.get('instructions') or None,
                'video_url': ex_video, 'row': row,
            })

    parsed_tracks = []
    for track in order:
        numbers = sorted(track['weeks'])
        first = track['first_week'] if track['first_week'] is not None else numbers[0]
        if numbers[0] < first:
            problems.error(track['row'], f'Track "{track["name"]}" has week {numbers[0]}, before its first week ({first}).')
            continue
        gaps = [n for n in range(first, numbers[-1] + 1) if n not in track['weeks']]
        if gaps:
            problems.error(track['row'], f'Track "{track["name"]}" has no rows for week(s) '
                                         f'{", ".join(map(str, gaps))}. Add a row with week_type "rest" for a rest week.')
        weeks = [track['weeks'][n] for n in numbers]
        for week in weeks:
            week['kind'] = week.get('kind') or 'training'
            if week['kind'] != 'rest' and not week['sessions']:
                problems.warn(week['row'], f'Track "{track["name"]}" week {week["number"]} has no sessions.')
            for s in week['sessions']:
                s['requirement'] = s.get('requirement') or 'required'
                s['session_type'] = s.get('session_type') or 'custom'
                if s['requirement'] == 'alternative' and not s.get('alternative_group'):
                    problems.error(s['row'], f'Session "{s["title"]}" is an alternative but has no alternative_group.')
        parsed_tracks.append({'name': track['name'], 'first_week': first, 'offset': track['offset'] or 0,
                              'weeks': weeks})
    if not parsed_tracks and not problems.errors:
        problems.error(1, 'The programme sheet has no rows.')

    tests, milestones = _parse_tests(test_rows or [], parsed_tracks, problems)
    return {'tracks': parsed_tracks, 'tests': tests, 'milestones': milestones,
            'errors': problems.errors, 'warnings': problems.warnings,
            'counts': {
                'tracks': len(parsed_tracks),
                'weeks': sum(len(t['weeks']) for t in parsed_tracks),
                'sessions': sum(len(w['sessions']) for t in parsed_tracks for w in t['weeks']),
                'exercises': sum(len(s['exercises']) for t in parsed_tracks for w in t['weeks'] for s in w['sessions']),
                'tests': len(tests), 'milestones': len(milestones),
            }}


def _parse_tests(rows, tracks, problems):
    records = _records(rows, TEST_COLUMNS, TEST_REQUIRED, problems, 'tests')
    by_track = {t['name'].lower(): t for t in tracks}
    tests, milestones = [], {}
    seen = set()
    for row, rec in records:
        label = f'Tests row {row}'
        name = rec.get('test')
        if not name:
            problems.error(row, f'{label}: no test name.')
            continue
        if name.lower() in seen:
            problems.error(row, f'{label}: "{name}" is listed twice.')
            continue
        seen.add(name.lower())
        if not rec.get('unit'):
            problems.error(row, f'{label}: "{name}" needs a unit.')
        measurement = (rec.get('measurement') or 'other').lower()
        if measurement not in MEASUREMENT_KEYS:
            problems.error(row, f'{label}: measurement must be one of {", ".join(sorted(MEASUREMENT_KEYS))}.')
        aggregation = (rec.get('aggregation') or 'best').lower()
        if aggregation not in AGGREGATION_KEYS:
            problems.error(row, f'{label}: aggregation must be best, mean or last.')
        better = (rec.get('better') or '').lower()
        if better not in ('higher', 'lower'):
            problems.error(row, f'{label}: better must be "higher" or "lower".')
        minimum = _float(rec.get('min'), row, problems, 'min')
        maximum = _float(rec.get('max'), row, problems, 'max')
        if minimum is not None and maximum is not None and minimum >= maximum:
            problems.error(row, f'{label}: min must be below max.')
        test = {'name': name, 'unit': rec.get('unit') or '', 'measurement': measurement,
                'attempts': _int(rec.get('attempts'), row, problems, 'attempts', 1) or 1,
                'aggregation': aggregation, 'higher_is_better': better == 'higher',
                'min_value': minimum, 'max_value': maximum,
                'description': rec.get('description') or None, 'instructions': rec.get('instructions') or None}
        tests.append(test)

        weeks = []
        if rec.get('baseline_week'):
            weeks.append(('baseline', rec['baseline_week']))
        for part in (rec.get('retest_weeks') or '').replace(';', ',').split(','):
            if part.strip():
                weeks.append(('retest', part.strip()))
        if not weeks:
            continue
        track = by_track.get((rec.get('milestone_track') or '').lower())
        if track is None:
            problems.error(row, f'{label}: milestone_track "{rec.get("milestone_track") or ""}" isn\'t a track '
                                'in the programme sheet.')
            continue
        for kind, raw in weeks:
            number = _int(raw, row, problems, f'{kind} week')
            if number is None:
                continue
            if not any(w['number'] == number for w in track['weeks']):
                problems.error(row, f'{label}: {track["name"]} has no week {number}.')
                continue
            ms = milestones.setdefault((track['name'], number, kind), {
                'kind': kind, 'track': track['name'], 'week': number, 'tests': [],
                'name': f'{"Baseline tests" if kind == "baseline" else "Retest"} ({track["name"]} week {number})'})
            ms['tests'].append(name)
    return tests, sorted(milestones.values(), key=lambda m: (m['track'], m['week']))


# --------------------------------------------------------------------------
# Building
# --------------------------------------------------------------------------

def find_duplicate(team_id, digest):
    return Programme.query.filter_by(team_organization_id=team_id, import_hash=digest).first()


def build(parsed, team_id, user, name, description=None, start_date=None, schedule_mode='calendar',
          kind='programme', filename=None, digest=None):
    """A new programme with a draft holding the parsed content. Only call
    with a parse that has no errors."""
    if parsed['errors']:
        raise ValueError('Fix the errors before importing.')
    programme = Programme(team_organization_id=team_id, name=name, description=description, kind=kind,
                          schedule_mode=schedule_mode, start_date=start_date, source_kind='import',
                          source_name=filename, import_filename=filename, import_hash=digest,
                          created_by_id=user.id if user else None)
    db.session.add(programme)
    version = ProgrammeVersion(number=1, status='draft', team_organization_id=team_id)
    programme.versions.append(version)
    week_lookup = {}
    for t_pos, t in enumerate(parsed['tracks']):
        track = ProgrammeTrack(name=t['name'], position=t_pos, first_week_number=t['first_week'],
                               start_offset_weeks=t['offset'], team_organization_id=team_id)
        version.tracks.append(track)
        for w in t['weeks']:
            week = ProgrammeWeek(position=w['number'] - t['first_week'], title=w.get('title'), kind=w['kind'],
                                 objective=w.get('objective'), supporting_skills=w.get('skills'),
                                 team_organization_id=team_id)
            track.weeks.append(week)
            week_lookup[(t['name'], w['number'])] = week
            for s_pos, s in enumerate(w['sessions']):
                session = ProgrammeSession(
                    position=s_pos, title=s['title'], session_type=s['session_type'],
                    custom_type=s.get('custom_type'), requirement=s['requirement'],
                    alternative_group=s.get('alternative_group'), duration_minutes=s.get('duration_minutes'),
                    focus=s.get('focus'), warmup=s.get('warmup'), cooldown=s.get('cooldown'),
                    coaching_cues=s.get('coaching_cues'), video_url=s.get('video_url'),
                    log_fields=s.get('log_fields'), team_organization_id=team_id)
                week.sessions.append(session)
                for e_pos, e in enumerate(s['exercises']):
                    session.exercises.append(ProgrammeExercise(
                        position=e_pos, team_organization_id=team_id,
                        **{k: e[k] for k in ('name', 'block', 'sets', 'reps', 'distance_m', 'duration_seconds',
                                             'rest', 'intensity', 'tempo', 'instructions', 'video_url')}))
    tests = {}
    for pos, t in enumerate(parsed['tests']):
        test = PerformanceTest(position=pos, team_organization_id=team_id, **t)
        version.tests.append(test)
        tests[t['name']] = test
    db.session.flush()
    for pos, m in enumerate(parsed['milestones']):
        version.milestones.append(ProgrammeMilestone(
            position=pos, name=m['name'], kind=m['kind'], week=week_lookup[(m['track'], m['week'])],
            test_keys=[tests[n].key for n in m['tests']], team_organization_id=team_id))
    db.session.flush()
    return programme


# --------------------------------------------------------------------------
# Export (the same format)
# --------------------------------------------------------------------------

def _num(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def export_rows(version):
    rows = [PROGRAMME_COLUMNS]
    for track in version.tracks:
        first_row_of_track = True
        for week in track.weeks:
            week_cells = {
                'track': track.name, 'week': str(week.number), 'week_title': week.title or '',
                'week_type': week.kind if week.kind != 'training' else '', 'week_objective': week.objective or '',
                'week_skills': week.supporting_skills or '',
            }
            if first_row_of_track:
                week_cells['track_first_week'] = str(track.first_week_number)
                week_cells['track_start_offset'] = str(track.start_offset_weeks)
                first_row_of_track = False
            if not week.sessions:
                rows.append([week_cells.get(c, '') for c in PROGRAMME_COLUMNS])
                continue
            for session in week.sessions:
                cells = dict(week_cells)
                cells.update({
                    'session': session.title,
                    'session_type': session.custom_type if session.session_type == 'custom' and session.custom_type
                    else session.session_type,
                    'requirement': session.requirement, 'alternative_group': session.alternative_group or '',
                    'session_minutes': _num(session.duration_minutes), 'session_focus': session.focus or '',
                    'warmup': session.warmup or '', 'cooldown': session.cooldown or '',
                    'coaching_cues': session.coaching_cues or '', 'session_video': session.video_url or '',
                    'log_fields': ', '.join(session.log_fields or []),
                })
                exercises = session.exercises or [None]
                for ex in exercises:
                    row = dict(cells)
                    if ex is not None:
                        row.update({
                            'exercise': ex.name, 'block': ex.block or '', 'sets': ex.sets or '', 'reps': ex.reps or '',
                            'distance_m': _num(ex.distance_m), 'duration_s': _num(ex.duration_seconds),
                            'rest': ex.rest or '', 'intensity': ex.intensity or '', 'tempo': ex.tempo or '',
                            'instructions': ex.instructions or '', 'exercise_video': ex.video_url or '',
                        })
                    rows.append([row.get(c, '') for c in PROGRAMME_COLUMNS])
    return rows


def export_test_rows(version):
    rows = [TEST_COLUMNS]
    for test in version.tests:
        mss = [m for m in version.milestones if m.is_test and m.week and m.covers(test)]
        track = mss[0].week.track if mss else None
        mss = [m for m in mss if m.week.track_id == (track.id if track else None)]
        baseline = next((m for m in mss if m.kind == 'baseline'), None)
        retests = [m for m in mss if m.kind == 'retest']
        cells = {
            'test': test.name, 'unit': test.unit, 'measurement': test.measurement, 'attempts': str(test.attempts),
            'aggregation': test.aggregation, 'better': 'higher' if test.higher_is_better else 'lower',
            'min': _num(test.min_value), 'max': _num(test.max_value), 'description': test.description or '',
            'instructions': test.instructions or '', 'milestone_track': track.name if track else '',
            'baseline_week': str(baseline.week.number) if baseline else '',
            'retest_weeks': ', '.join(str(m.week.number) for m in retests),
        }
        rows.append([cells[c] for c in TEST_COLUMNS])
    return rows
